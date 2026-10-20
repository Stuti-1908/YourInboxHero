"""Read an uploaded CSV or Excel (.xlsx) file into header-keyed rows.

Both importers (invoices, document requests) share this so customers can
upload whatever their spreadsheet tool exports. Every value comes back as a
stripped string; Excel dates become YYYY-MM-DD and whole numbers lose their
trailing ".0", so downstream validation is identical for both formats.
"""
import csv
import io
import zipfile
from datetime import date, datetime
from typing import List, Tuple

from fastapi import HTTPException

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
# Importers cap rows at 1000; reading one past that lets them report the
# overflow without ever materialising a huge (or decompression-bomb) sheet.
MAX_DATA_ROWS = 1001
# An .xlsx is a zip; a small upload can expand to gigabytes (zip bomb).
# Real spreadsheets of 1000 rows unpack to a few MB.
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xlsx(raw: bytes) -> Tuple[List[str], List[dict]]:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            unpacked = sum(i.file_size for i in zf.infolist())
    except zipfile.BadZipFile:
        raise HTTPException(status_code=400, detail="Could not read the Excel file. Save it as .xlsx or CSV and try again.")
    if unpacked > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise HTTPException(status_code=400, detail="This Excel file is too large to import. Split it or save it as CSV.")

    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the Excel file. Save it as .xlsx or CSV and try again.")
    ws = wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise HTTPException(status_code=400, detail="File is empty")
    headers = [_cell_to_str(h).lower() for h in header_row]
    rows = []
    for row_number, values in enumerate(rows_iter, start=2):
        if values is None or all(v is None or str(v).strip() == "" for v in values):
            continue
        row = {headers[i]: _cell_to_str(v) for i, v in enumerate(values) if i < len(headers) and headers[i]}
        row["__row__"] = row_number
        rows.append(row)
        if len(rows) >= MAX_DATA_ROWS:
            break
    wb.close()
    return headers, rows


def _read_csv(raw: bytes) -> Tuple[List[str], List[dict]]:
    try:
        text = raw.decode("utf-8-sig")  # handles Excel's BOM-prefixed UTF-8 exports
    except UnicodeDecodeError:
        raise HTTPException(status_code=400, detail="CSV file must be UTF-8 encoded")
    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        raise HTTPException(status_code=400, detail="File is empty")
    headers = [h.strip().lower() for h in reader.fieldnames]
    rows = []
    for raw_row in reader:
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in raw_row.items() if k is not None}
        if not any(row.values()):
            continue
        # line_num is the physical line just read, so it stays right across
        # skipped blank lines (multi-line quoted cells report their last line).
        row["__row__"] = reader.line_num
        rows.append(row)
        if len(rows) >= MAX_DATA_ROWS:
            break
    return headers, rows


def read_table(raw: bytes, filename: str) -> Tuple[List[str], List[dict]]:
    """Return (lowercased headers, rows). Rows are in file order and each
    carries its spreadsheet row number under "__row__" (header = row 1), so
    error messages point at the right line even when blank rows are skipped."""
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=400, detail="File is larger than 5 MB")
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(raw)
    if name.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Old .xls files aren't supported. Save as .xlsx or CSV.")
    return _read_csv(raw)
